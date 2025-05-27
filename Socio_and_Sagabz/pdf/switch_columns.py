from docx import Document
import os
import glob

def switch_columns_in_table(doc_path, col1_title, col2_title):
    # Load the Word document
    doc = Document(doc_path)
    
    for table in doc.tables:
        # Find the column indices by their titles
        col1_index = col2_index = None
        for i, cell in enumerate(table.rows[0].cells):
            if cell.text.strip() == col1_title:
                col1_index = i
            elif cell.text.strip() == col2_title:
                col2_index = i
        
        # If both columns are found, switch their values
        if col1_index is not None and col2_index is not None:
            for row in table.rows[1:]:  # Skip the header row
                col1_cell = row.cells[col1_index]
                col2_cell = row.cells[col2_index]

                # Copy the runs (styled text) from one cell to another
                col1_runs = col1_cell.paragraphs[0].runs
                col2_runs = col2_cell.paragraphs[0].runs

                # Store the runs temporarily
                col1_text = [(run.text, run.bold, run.italic, run.underline, run.font.size) for run in col1_runs]
                col2_text = [(run.text, run.bold, run.italic, run.underline, run.font.size) for run in col2_runs]

                # Clear the paragraphs
                col1_cell.paragraphs[0].clear()
                col2_cell.paragraphs[0].clear()

                # Write the runs back to the opposite cells
                for text, bold, italic, underline, size in col2_text:
                    run = col1_cell.paragraphs[0].add_run(text)
                    run.bold = bold
                    run.italic = italic
                    run.underline = underline
                    run.font.size = size
                    run.font.name = "David"

                for text, bold, italic, underline, size in col1_text:
                    run = col2_cell.paragraphs[0].add_run(text)
                    run.bold = bold
                    run.italic = italic
                    run.underline = underline
                    run.font.size = size
                    run.font.name = "David"

    # Save the modified document
    doc.save(doc_path)
    

if __name__ == "__main__":
    
    for file in glob.glob("input/*.docx"):
        switch_columns_in_table(file, "סמסטר א׳", "סמסטר ב׳")
    
